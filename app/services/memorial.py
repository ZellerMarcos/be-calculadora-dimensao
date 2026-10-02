from __future__ import annotations

from io import BytesIO
import re
import unicodedata

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


def _safe_text(value: object) -> str:
    text = unicodedata.normalize("NFKC", str(value))
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)
    return text[:2000]


def _add_key_value(document: Document, label: str, value: object) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.add_run(f"{label}: ").bold = True
    paragraph.add_run(_safe_text(value) or "Não informado")


def build_docx(memorial: dict) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(9)
    styles["Heading 1"].font.color.rgb = RGBColor(18, 55, 46)
    styles["Heading 2"].font.color.rgb = RGBColor(8, 127, 125)

    title = document.add_heading("Memorial de cálculo de calha pluvial", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph("Dimensionamento e verificação hidráulica")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].italic = True
    _add_key_value(document, "Norma", memorial.get("norma", "ABNT NBR 10844:1989"))
    _add_key_value(document, "Versão do motor", memorial.get("versaoMotor", ""))
    _add_key_value(document, "Gerado em", memorial.get("geradoEm", ""))

    document.add_heading("2. Dados de entrada e premissas", level=1)
    inputs = memorial.get("entradas", {})
    input_labels = (
        ("station_name", "Posto pluviométrico"),
        ("return_period", "Período de retorno (anos)"),
        ("intensity", "Intensidade (mm/h)"),
        ("roof_width", "Projeção horizontal a (m)"),
        ("roof_rise", "Altura h (m)"),
        ("roof_surface", "Tipo de cobertura"),
        ("gutter_length", "Extensão da calha (m)"),
        ("outlet_type", "Posição da saída"),
        ("extension_side_a", "Extensão do lado A (m)"),
        ("extension_side_b", "Extensão do lado B (m)"),
        ("profile", "Seção"),
        ("bottom_width_mm", "Largura de fundo / diâmetro (mm)"),
        ("top_width_mm", "Largura superior (mm)"),
        ("useful_depth_mm", "Lâmina informada (mm)"),
        ("freeboard_mm", "Bordo livre adotado (mm)"),
        ("slope_percent", "Declividade (%)"),
        ("roughness", "Coeficiente de rugosidade n"),
        ("material", "Material"),
        ("curve_condition", "Condição de mudança de direção"),
    )
    for key, label in input_labels:
        _add_key_value(document, label, inputs.get(key, ""))

    document.add_heading("Premissas e critérios", level=2)
    for premise in memorial.get("premissas", []):
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(f"[{_safe_text(premise.get('origem', ''))}] ").bold = True
        paragraph.add_run(_safe_text(premise.get("descricao", "")))

    document.add_heading("3. Desenvolvimento dos cálculos", level=1)
    for step in memorial.get("passos", []):
        document.add_heading(f"{step.get('ordem', '')}. {_safe_text(step.get('titulo', ''))}", level=2)
        _add_key_value(document, "Referência normativa", step.get("referenciaNorma", ""))
        _add_key_value(document, "Fórmula", step.get("formulaTexto", ""))
        _add_key_value(document, "Expressão", step.get("formulaLatex", ""))
        _add_key_value(document, "Substituição numérica", step.get("substituicao", ""))
        unit = step.get("unidade", "")
        _add_key_value(document, "Resultado", f"{step.get('resultado', '')} {unit}".strip())

    document.add_heading("4. Resultados e dimensionamento", level=1)
    results = memorial.get("resultados", {})
    for key, label in (
        ("areaContribuicaoM2", "Área de contribuição (m²)"),
        ("coeficienteTabela1", "Coeficiente da Tabela 1"),
        ("vazaoBaseLmin", "Vazão sem correção (L/min)"),
        ("vazaoProjetoLmin", "Vazão de projeto (L/min)"),
        ("areaMolhadaM2", "Área molhada (m²)"),
        ("perimetroMolhadoM", "Perímetro molhado (m)"),
        ("raioHidraulicoM", "Raio hidráulico (m)"),
        ("vazaoCapacidadeLmin", "Capacidade (L/min)"),
        ("taxaAtendimento", "Taxa de atendimento"),
        ("velocidadeMs", "Velocidade informativa (m/s)"),
        ("alturaTotalMm", "Altura total (mm)"),
    ):
        _add_key_value(document, label, results.get(key, ""))

    dimensioning = memorial.get("dimensionamento")
    if dimensioning:
        document.add_heading("Dimensionamento mínimo calculado", level=2)
        for key, label in (
            ("status", "Situação do dimensionamento"),
            ("secao", "Seção"),
            ("larguraFundoMm", "Largura de fundo (mm)"),
            ("taludeZ", "Talude z"),
            ("diametroMm", "Diâmetro interno (mm)"),
            ("laminaMm", "Lâmina dimensionada (mm)"),
            ("passoConstrutivoMm", "Passo construtivo (mm)"),
            ("capacidadeLmin", "Capacidade após dimensionamento (L/min)"),
        ):
            _add_key_value(document, label, dimensioning.get(key, ""))

    document.add_heading("5. Verificações e conclusão", level=1)
    alerts = memorial.get("alertas", [])
    if alerts:
        for alert in alerts:
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.add_run(f"{_safe_text(alert.get('nivel', '').upper())} · {_safe_text(alert.get('codigo', ''))}: ").bold = True
            paragraph.add_run(_safe_text(alert.get("mensagem", "")))
            paragraph.add_run(f" (Ref.: {_safe_text(alert.get('referenciaNorma', ''))})")
    else:
        document.add_paragraph("Nenhum alerta adicional para as entradas avaliadas.")

    conclusion = document.add_paragraph(_safe_text(memorial.get("conclusao", "")))
    conclusion.runs[0].bold = True
    document.add_paragraph(
        "Este memorial registra o cálculo efetuado a partir das entradas informadas. "
        "A análise e a responsabilidade técnica permanecem com o profissional habilitado."
    )

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()